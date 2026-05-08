#!/usr/bin/env python3
r"""
WES2 spec-grounded reviewer/implementer orchestrator.

Purpose:
- Claude Code implements WES2 in small phases.
- OpenAI reviews each phase against a compact WES2 spec brief + relevant excerpts.
- The script avoids sending the full spec on every review call, reducing TPM failures.
- The script enforces a WES2-only write scope.

Run from repo root.

Example:
python .\docs\wes2_orchestrator.py --spec ".\docs\wes2\WES2_Product_Spec_and_Implementation_Contract_v1_revised_with_examples.docx" --start-phase 18 --max-phases 1
"""

from __future__ import annotations

import argparse
import fnmatch
import json
import os
import re
import shlex
import subprocess
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

try:
    from openai import OpenAI
except Exception:
    OpenAI = None


# ─────────────────────────────────────────────────────────────────────────────
# Write-scope policy
# ─────────────────────────────────────────────────────────────────────────────

ALLOWED_PREFIXES = [
    "docs/wes2/",
    ".orchestrator/wes2/",
]

ALLOWED_GLOBS = [
    "lib/WES2_*.dart",
    "lib/WES2_widgets/WES2_*.dart",
    "test/wes2/WES2_*.dart",
]

ALLOWED_EXACT = [
    # Only for the original WES2 quick-access card/import. Later phases normally
    # should not need this, but the allow-list keeps the script reusable.
    "lib/home_screen.dart",
]

BLOCKED_GLOBS = [
    "lib/workout_entry_screen.dart",
    "lib/main.dart",
    "lib/bb3_*.dart",
    "lib/periodization_model_utils.dart",
    "lib/progression_engine.dart",
    "pubspec.yaml",
    "assets/**",
    "**/*.g.dart",
    "**/*.freezed.dart",
    "**/*.mocks.dart",
    ".claude/settings.local.json",
]


PHASE_MAP = {
    17: {
        "name": "Day Timer + Workout Summary",
        "keywords": [
            "timer",
            "summary",
            "workout summary",
            "floating timer",
            "sets logged",
            "Add Circuit",
            "session time",
        ],
    },
    18: {
        "name": "Templates",
        "keywords": [
            "template",
            "templates",
            "save as template",
            "load template",
            "template-loaded",
            "templateLoaded",
        ],
    },
    19: {
        "name": "Settings Cog Dialog",
        "keywords": [
            "settings",
            "exerciseSettings",
            "increments",
            "weekly frequency",
            "periodization",
            "RIR",
            "progression",
        ],
    },
    20: {
        "name": "BB3 Structural Handoff",
        "keywords": [
            "BB3",
            "planned day",
            "delete",
            "replace",
            "move",
            "structural",
            "updatePlannedDay",
        ],
    },
    21: {
        "name": "Hint Engine Integration",
        "keywords": [
            "hint",
            "hints",
            "E1RM",
            "PMU",
            "progression",
            "ProgressionEngine",
            "planned override",
        ],
    },
}


@dataclass
class CmdResult:
    code: int
    out: str


def run(cmd: list[str] | str, *, timeout: int = 1200, input_text: str | None = None) -> CmdResult:
    shell = isinstance(cmd, str)
    p = subprocess.run(
        cmd,
        input=input_text,
        text=True,
        shell=shell,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        timeout=timeout,
    )
    return CmdResult(p.returncode, p.stdout)


def die(msg: str, code: int = 1) -> None:
    print(f"\n❌ {msg}\n")
    sys.exit(code)


def repo_root() -> Path:
    r = run(["git", "rev-parse", "--show-toplevel"], timeout=30)
    if r.code != 0:
        die("This script must be run inside a git repo.")
    return Path(r.out.strip())


def _extract_docx_table(table) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return rows


def read_spec(path: Path) -> str:
    if not path.exists():
        die(f"Spec not found: {path}")

    if path.suffix.lower() == ".docx":
        try:
            import docx
        except Exception:
            die("Reading .docx requires: pip install python-docx")
        d = docx.Document(str(path))
        parts: list[str] = []
        for p in d.paragraphs:
            txt = p.text.strip()
            if txt:
                parts.append(txt)
        for t in d.tables:
            parts.extend(_extract_docx_table(t))
        return "\n".join(parts)

    return path.read_text(encoding="utf-8", errors="replace")


def ensure_clean_or_confirm() -> None:
    r = run(["git", "status", "--porcelain"], timeout=30)
    if r.code != 0:
        die("Could not read git status.")
    if r.out.strip():
        die(
            "Working tree is not clean. Commit/stash current work first so scope enforcement is safe.\n\n"
            "Run: git status --short"
        )


def changed_files() -> list[str]:
    r = run(["git", "diff", "--name-only"], timeout=60)
    if r.code != 0:
        die("Could not read git diff.")
    files = [x.strip().replace("\\", "/") for x in r.out.splitlines() if x.strip()]

    r2 = run(["git", "ls-files", "--others", "--exclude-standard"], timeout=60)
    if r2.code == 0:
        files += [x.strip().replace("\\", "/") for x in r2.out.splitlines() if x.strip()]

    return sorted(set(files))


def matches_any(path: str, patterns: Iterable[str]) -> bool:
    p = path.replace("\\", "/")
    return any(fnmatch.fnmatch(p, pat) for pat in patterns)


def allowed(path: str) -> bool:
    p = path.replace("\\", "/")
    if p in ALLOWED_EXACT:
        return True
    if any(p.startswith(pref) for pref in ALLOWED_PREFIXES):
        return True
    if matches_any(p, ALLOWED_GLOBS):
        return True
    return False


def blocked(path: str) -> bool:
    return matches_any(path.replace("\\", "/"), BLOCKED_GLOBS)


def enforce_scope(*, revert: bool) -> tuple[list[str], list[str]]:
    cf = changed_files()
    out_of_scope = [p for p in cf if not allowed(p)]
    blocked_touched = [p for p in cf if blocked(p)]
    bad = sorted(set(out_of_scope + blocked_touched))

    if bad and revert:
        for p in bad:
            tracked = run(["git", "ls-files", "--error-unmatch", p], timeout=30)
            if tracked.code == 0:
                run(["git", "checkout", "--", p], timeout=60)
            else:
                pp = Path(p)
                try:
                    if pp.is_file() or pp.is_symlink():
                        pp.unlink()
                except FileNotFoundError:
                    pass
    return out_of_scope, blocked_touched


def get_allowed_changed_files() -> list[str]:
    return [p for p in changed_files() if allowed(p)]


def git_diff_allowed(max_chars: int = 45000) -> str:
    files = get_allowed_changed_files()
    if not files:
        return ""

    tracked_files = [p for p in files if run(["git", "ls-files", "--error-unmatch", p], timeout=10).code == 0]
    untracked = [p for p in files if p not in tracked_files]
    chunks: list[str] = []

    if tracked_files:
        r = run(["git", "diff", "--"] + tracked_files, timeout=300)
        chunks.append(r.out)

    for p in untracked:
        try:
            txt = Path(p).read_text(encoding="utf-8", errors="replace")
            chunks.append(f"\n--- UNTRACKED FILE: {p} ---\n{txt[:20000]}")
        except Exception:
            chunks.append(f"\n--- UNTRACKED FILE: {p} (binary/unreadable) ---\n")

    txt = "\n".join(chunks)
    return txt[-max_chars:]


def write_log(log_dir: Path, name: str, text: str) -> Path:
    log_dir.mkdir(parents=True, exist_ok=True)
    p = log_dir / name
    p.write_text(text, encoding="utf-8")
    return p


def claude(prompt: str, *, timeout: int, claude_cmd: str) -> str:
    cmd = shlex.split(claude_cmd)
    r = run(cmd, timeout=timeout, input_text=prompt)
    if r.code != 0:
        die(f"Claude command failed:\n{r.out}")
    return r.out


def openai_review(prompt: str, *, model: str, max_output_tokens: int = 2200) -> dict:
    if OpenAI is None:
        die("OpenAI SDK missing. Install with: pip install openai")
    client = OpenAI()
    resp = client.responses.create(
        model=model,
        input=prompt,
        text={"format": {"type": "json_object"}},
        max_output_tokens=max_output_tokens,
    )
    text = resp.output_text
    try:
        return json.loads(text)
    except Exception:
        return {"approved": False, "reason": "Reviewer did not return valid JSON.", "raw": text}


def scope_text() -> str:
    return """
Allowed writes:
- lib/WES2_*.dart
  Examples explicitly allowed:
  - lib/WES2_screen.dart
  - lib/WES2_controller.dart
  - lib/WES2_models.dart
  - lib/WES2_repository.dart
  - lib/WES2_plan_service.dart
  - lib/WES2_local_store.dart
  - lib/WES2_template_service.dart
  - lib/WES2_sync_service.dart
- lib/WES2_widgets/WES2_*.dart
  Examples explicitly allowed:
  - lib/WES2_widgets/WES2_day_actions_row.dart
  - lib/WES2_widgets/WES2_set_row.dart
  - lib/WES2_widgets/WES2_exercise_card.dart
  - lib/WES2_widgets/WES2_exercise_picker.dart
- test/wes2/WES2_*.dart
- docs/wes2/**
- .orchestrator/wes2/**
- lib/home_screen.dart ONLY for adding/updating the WES2 Quick Access card/import

Read-only inspection:
- Existing non-WES2 files may be read for schema/reference only when needed.
- Reading lib/templates.dart, lib/template_model.dart, BB3 files, current WES, PMU, or progression files is allowed when necessary for schema/behaviour reference.
- Writing to those files is NOT allowed unless they match the allowed write scope above.

Blocked writes:
- lib/workout_entry_screen.dart
- lib/main.dart
- lib/bb3_*.dart
- lib/periodization_model_utils.dart
- lib/progression_engine.dart
- pubspec.yaml
- assets/**
- generated files
- .claude/settings.local.json

Important distinction:
- Allowed write scope controls what Claude may modify.
- Read-only inspection of existing files is allowed when necessary to understand schema or behaviour.
- The reviewer must not reject a plan merely because it proposes reading a blocked/reference file.
- The reviewer must reject if the plan proposes editing a blocked/reference file.

Do not rename, move, delete, or modify asset files or asset references.
Do not use broad scans.
Do not use ListTile inside PopupMenuItem.
Avoid Flutter RenderFlex overflow: long Row text must use Expanded/Flexible, maxLines, and ellipsis.
Treat any RenderFlex overflow as a blocker.
""".strip()


# ─────────────────────────────────────────────────────────────────────────────
# Spec compression/excerpts
# ─────────────────────────────────────────────────────────────────────────────

def build_spec_brief(spec: str) -> str:
    keep_patterns = [
        r"(?is)1\. Core Contract.*?(?=\n2\. Existing|\Z)",
        r"(?is)3\. Firestore and Data Model Contract.*?(?=\n4\. State|\Z)",
        r"(?is)4\. State Ownership.*?(?=\n5\. Opening|\Z)",
        r"(?is)5\. Opening, Loading.*?(?=\n6\. Planned|\Z)",
        r"(?is)9\..*?Exercise, set, circuit, template, undo, and delete/replace behavior.*?(?=\n10\.|\Z)",
        r"(?is)10\..*?UI specification.*?(?=\n11\.|\Z)",
        r"(?is)11\..*?Timed exercises, velocity, notes, video, timer, and summary behavior.*?(?=\n12\.|\Z)",
        r"(?is)12\..*?Multi-device.*?(?=\n13\.|\Z)",
        r"(?is)Acceptance criteria and test matrix.*?\Z",
    ]

    chunks: list[str] = []
    for pat in keep_patterns:
        m = re.search(pat, spec)
        if m:
            chunks.append(m.group(0)[:5000])

    brief = "\n\n--- SPEC BRIEF SECTION ---\n\n".join(chunks)
    if len(brief) < 2000:
        brief = spec[:18000]
    return brief[:24000]


def relevant_spec_excerpt(spec: str, phase_no: int, max_chars: int = 7000) -> str:
    keywords = PHASE_MAP.get(phase_no, {}).get("keywords", [])
    if not keywords:
        return spec[:max_chars]

    lines = spec.splitlines()
    hits: list[int] = []
    for i, line in enumerate(lines):
        low = line.lower()
        if any(k.lower() in low for k in keywords):
            hits.append(i)

    windows: list[str] = []
    seen_ranges: list[tuple[int, int]] = []
    for i in hits[:18]:
        start = max(0, i - 18)
        end = min(len(lines), i + 45)
        if any(not (end < a or start > b) for a, b in seen_ranges):
            continue
        seen_ranges.append((start, end))
        windows.append("\n".join(lines[start:end]))

    txt = "\n\n--- RELEVANT SPEC EXCERPT ---\n\n".join(windows)
    if not txt.strip():
        txt = build_spec_brief(spec)

    return txt[:max_chars]


def phase_guidance(phase_no: int) -> str:
    if phase_no == 17:
        return """
Phase 17 target:
- Day-level floating timer and Workout Summary.
- Timer AppBar overflow item should open/toggle floating timer.
- Templates remains placeholder.
- Summary should be accessible from the bottom/Add Circuit row if feasible because the spec places Summary there.
- Do not implement timed-exercise reps-as-seconds.
- Do not persist timer to Firestore/Isar.
- Do not touch hints/PMU/progression/templates/settings.
""".strip()

    if phase_no == 18:
        return """
Phase 18 target:
- Templates.
- Read existing template/schema reference files only as needed; do not edit them.
- Do not guess template schema.
- Keep all writes inside WES2-specific files.
- Loading a template should not corrupt BB3 planned data unless the spec path and write shape are understood and the write target is a WES2-allowed file/method.
- Preserve duplicate exercise prevention by exerciseId.
- Avoid broad controller/repository rewrites.
""".strip()

    if phase_no == 19:
        return """
Phase 19 target:
- Settings cog dialog.
- Must use block-level exerciseSettings only; do not write plannedExerciseDetails.
- Requires activeBlockId guard.
- Internet-required; no offline queue.
- Do not implement hint recalculation beyond reload/refresh trigger.
""".strip()

    if phase_no == 20:
        return """
Phase 20 target:
- BB3 structural handoff for BB3-sourced delete/replace/move/reorder.
- Read planned day write path carefully.
- Do not mutate BB3 until path and row shape are verified.
""".strip()

    if phase_no == 21:
        return """
Phase 21 target:
- Hint engine integration.
- Must respect BB3 planned overrides as hints only, never actual values.
- Do not overwrite typed WES2 values.
""".strip()

    return "Continue the next smallest safe WES2 phase from the approved plan."


# ─────────────────────────────────────────────────────────────────────────────
# Prompts
# ─────────────────────────────────────────────────────────────────────────────

def initial_plan_prompt(spec_brief: str, start_phase: int) -> str:
    return f"""
You are Claude Code implementing WES2 for the GOODLIFT Flutter/Firebase app.

Do NOT edit files yet.
Produce a concise continuation plan only, then STOP.

We are not starting WES2 from scratch.
Continue from current repo state.
Start at phase {start_phase}: {PHASE_MAP.get(start_phase, {}).get("name", "next WES2 phase")}.

{scope_text()}

Required plan:
1. Confirm current phase target.
2. Files/anchors to inspect.
3. Small implementation steps.
4. Spec rules you will follow by section/topic name.
5. Analyzer command.
6. Risks and overflow guards.
7. Explicitly distinguish READ-only reference files from writable WES2 files.

COMPACT SPEC BRIEF:
{spec_brief}
"""


def reviewer_plan_prompt(spec_brief: str, plan: str, start_phase: int) -> str:
    return f"""
You are the strict WES2 reviewer.

Return JSON only:
{{
  "approved": true/false,
  "reason": "...",
  "required_changes": ["..."],
  "scope_risks": ["..."]
}}

Review the continuation plan against the compact spec brief and write scope.
The plan must continue from phase {start_phase}, not restart WES2.

Must enforce:
{scope_text()}

Critical reviewer interpretation:
- lib/WES2_screen.dart, lib/WES2_controller.dart, lib/WES2_models.dart, lib/WES2_repository.dart, lib/WES2_plan_service.dart, lib/WES2_local_store.dart, lib/WES2_template_service.dart, and lib/WES2_sync_service.dart are allowed write targets because they match lib/WES2_*.dart.
- lib/WES2_widgets/WES2_*.dart files are allowed write targets.
- Do NOT reject the plan for READ-only inspection of reference/schema files such as lib/templates.dart or lib/template_model.dart.
- Reject only if the plan proposes WRITING to blocked/reference files.

Reviewer rule:
If you request a correction, cite/spec-reference the spec topic or section name that supports it.

COMPACT SPEC BRIEF:
{spec_brief}

PLAN:
{plan}
"""


def implement_prompt(spec_excerpt: str, plan: str, feedback: str, phase_no: int) -> str:
    return f"""
You are Claude Code implementing WES2 phase {phase_no}: {PHASE_MAP.get(phase_no, {}).get("name", "next phase")}.

Implement ONLY this phase.
Keep it small and reviewable.

{scope_text()}

Phase-specific guidance:
{phase_guidance(phase_no)}

Hard restrictions:
- Do not edit blocked files.
- If a change requires editing a blocked file, print BLOCKED with exact reason.
- Read-only inspection of reference/schema files is allowed when needed.
- Do not touch hints/PMU/progression unless this is explicitly the hint phase.
- Do not rename, move, delete, or modify asset files or asset references.
- Avoid RenderFlex overflows.

At the end, print:
PHASE_DONE: <summary>
or
BLOCKED: <reason>

Reviewer feedback to respect:
{feedback}

Approved continuation plan:
{plan[:8000]}

Relevant spec excerpt for this phase:
{spec_excerpt}
"""


def reviewer_diff_prompt(
    spec_excerpt: str,
    plan: str,
    diff: str,
    analyzer: str,
    out_of_scope: list[str],
    blocked_touched: list[str],
    phase_output: str,
    phase_no: int,
) -> str:
    return f"""
You are the strict WES2 reviewer.

Return JSON only:
{{
  "approved": true/false,
  "done": true/false,
  "reason": "...",
  "required_changes": ["..."],
  "spec_references_used": ["..."],
  "scope_violation": true/false,
  "next_phase_guidance": "..."
}}

Approve only if:
- changes stay within allowed WES2 write scope,
- no blocked file was edited,
- read-only inspection of reference files is not a scope violation; only writes are scope violations,
- implementation follows the relevant spec excerpt,
- analyzer output is acceptable or only unrelated pre-existing warnings,
- no RenderFlex/overflow-prone layout was introduced,
- the phase remains small and reviewable.

Critical reviewer interpretation:
- lib/WES2_screen.dart, lib/WES2_controller.dart, lib/WES2_models.dart, lib/WES2_repository.dart, lib/WES2_plan_service.dart, lib/WES2_local_store.dart, lib/WES2_template_service.dart, and lib/WES2_sync_service.dart are allowed write targets because they match lib/WES2_*.dart.
- lib/WES2_widgets/WES2_*.dart files are allowed write targets.
- Do not flag read-only inspection of non-WES2 reference files as a violation.

Reviewer rule:
If you request a correction, include the spec topic/section in spec_references_used and explain the correction.

Phase {phase_no}: {PHASE_MAP.get(phase_no, {}).get("name", "next phase")}

Allowed write scope:
{scope_text()}

Out-of-scope files detected by script:
{out_of_scope}

Blocked files detected by script:
{blocked_touched}

Claude phase output:
{phase_output[-8000:]}

Analyzer/test output:
{analyzer[-7000:]}

Diff:
{diff[-35000:]}

Relevant spec excerpt:
{spec_excerpt}

Approved continuation plan:
{plan[:5000]}
"""


def run_analyzer() -> str:
    parts: list[str] = []
    files = [p for p in get_allowed_changed_files() if p.endswith(".dart")]
    if files:
        fmt = run(["dart", "format"] + files, timeout=300)
        parts.append("== dart format ==\n" + fmt.out)
        ana = run(["flutter", "analyze"] + files, timeout=1200)
        parts.append("== targeted flutter analyze ==\n" + ana.out)
    else:
        parts.append("== dart format ==\n(no changed Dart files)")
        parts.append("== targeted flutter analyze ==\n(no changed Dart files)")
    return "\n\n".join(parts)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--spec", required=True)
    ap.add_argument("--max-phases", type=int, default=1)
    ap.add_argument("--start-phase", type=int, default=18)
    ap.add_argument("--claude-cmd", default=os.environ.get("CLAUDE_CMD", "claude -p"))
    ap.add_argument("--openai-model", default=os.environ.get("OPENAI_REVIEW_MODEL"))
    ap.add_argument("--revert-out-of-scope", action=argparse.BooleanOptionalAction, default=True)
    args = ap.parse_args()

    if not args.openai_model:
        die("Set OPENAI_REVIEW_MODEL, e.g. $env:OPENAI_REVIEW_MODEL='gpt-5.4-mini'")

    root = repo_root()
    os.chdir(root)
    ensure_clean_or_confirm()

    log_dir = root / ".orchestrator" / "wes2" / time.strftime("%Y%m%d_%H%M%S")
    spec = read_spec(Path(args.spec))
    if len(spec) < 10000:
        die("Spec text looks too short. Check the --spec path or DOCX extraction.")

    spec_brief = build_spec_brief(spec)
    write_log(log_dir, "00_spec_brief.txt", spec_brief)

    print("→ Asking Claude for continuation plan...")
    plan = claude(initial_plan_prompt(spec_brief, args.start_phase), timeout=1800, claude_cmd=args.claude_cmd)
    write_log(log_dir, "01_claude_plan.txt", plan)

    print("→ Reviewing continuation plan with OpenAI...")
    plan_review = openai_review(reviewer_plan_prompt(spec_brief, plan, args.start_phase), model=args.openai_model)
    write_log(log_dir, "02_plan_review.json", json.dumps(plan_review, indent=2))

    if not plan_review.get("approved"):
        die("Plan not approved. See log: " + str(log_dir / "02_plan_review.json"))

    feedback = json.dumps(plan_review, indent=2)

    for offset in range(args.max_phases):
        phase_no = args.start_phase + offset
        spec_excerpt = relevant_spec_excerpt(spec, phase_no)
        write_log(log_dir, f"phase_{phase_no:02d}_spec_excerpt.txt", spec_excerpt)

        print(f"\n→ Claude implementing phase {phase_no}: {PHASE_MAP.get(phase_no, {}).get('name', 'next phase')}...")
        out = claude(
            implement_prompt(spec_excerpt, plan, feedback, phase_no),
            timeout=3600,
            claude_cmd=args.claude_cmd,
        )
        write_log(log_dir, f"phase_{phase_no:02d}_claude.txt", out)

        out_of_scope, blocked_touched = enforce_scope(revert=args.revert_out_of_scope)
        if out_of_scope or blocked_touched:
            print("⚠️ Scope issue detected and reverted/removed where possible:")
            for b in sorted(set(out_of_scope + blocked_touched)):
                print(" -", b)

        print("→ Running formatter/analyzer...")
        analyzer = run_analyzer()
        write_log(log_dir, f"phase_{phase_no:02d}_analyzer.txt", analyzer)

        diff = git_diff_allowed()
        write_log(log_dir, f"phase_{phase_no:02d}_diff.patch", diff)

        print("→ Reviewing phase with OpenAI...")
        review = openai_review(
            reviewer_diff_prompt(
                spec_excerpt,
                plan,
                diff,
                analyzer,
                out_of_scope,
                blocked_touched,
                out,
                phase_no,
            ),
            model=args.openai_model,
        )
        write_log(log_dir, f"phase_{phase_no:02d}_review.json", json.dumps(review, indent=2))

        feedback = json.dumps(review, indent=2)

        if not review.get("approved"):
            print("Phase not approved. Claude will receive reviewer feedback next phase.")
            continue

        print(f"✅ Phase {phase_no} approved.")

        if review.get("done") and phase_no >= 21:
            print("\n✅ WES2 orchestration complete.")
            print("Logs:", log_dir)
            return

    print("\nReached requested phase count.")
    print("Logs:", log_dir)
    print("Review final changed files with: git diff --stat")


if __name__ == "__main__":
    main()
